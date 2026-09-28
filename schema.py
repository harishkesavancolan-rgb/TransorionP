"""
schema.py
---------
Builds a JSON schema (per sheet) directly from the column headers of the
customs export/import template (.xlsx), so the extractor never needs the
column names hardcoded.

Honest caveat: the column *headers themselves* are read dynamically from
whichever template you point at, but the aliases and type-forcing below
(_MANUAL_ALIASES, _STRING_FORCE) are hand-curated India-customs domain
knowledge tied to specific column names (e.g. "FTA Code" is always a code,
never a number; "Item_RITC" is commonly called "HSN Code" on invoices).
That part is intentionally hardcoded and has to stay that way -- customs
data needs it. If you rename a column in the template, the schema will
still pick up the new header automatically, but any alias/force-type rule
written against the *old* name will silently stop applying to it. There's
no automatic detection of that today; re-check this file's lists after
renaming columns.

Also generates human-readable aliases for each column name, so the LLM
prompt can show the model what common invoice labels map to which schema
field.  For example:  Item_Country_Org  ->  "Country of Origin".
"""
from __future__ import annotations
from pathlib import Path
from typing import Any

import openpyxl

# Heuristics for inferring a field's data type purely from its column name.
_NUMBER_HINTS = ("qty", "price", "amt", "amount", "percentage", "value",
                  "val", "rate", "duty", "pmv", "pkg", "taxable")
_DATE_HINTS = ("date",)

# Fields that are ALWAYS strings despite potentially matching number hints.
# Y/N flags, codes, text fields, etc.
_STRING_FORCE = {
    "item_cess", "item_accessory", "item_thirdparty", "item_quota",
    "item_ar4", "item_reward", "item_str", "item_jnoti_no",
    "per", "printed", "itm_source_cntry", "itm_transit_cntry",
    "itmigststatus", "item_end_use", "fta code", "state code",
    "district code", "cess per", "cess rate", "cess amount",
    "item_scheme_code", "itm_manu_codetype", "accessory_status",
    "item_ser_no", "sl_no",
}

# No system fields excluded — all fields including "Printed" are real.
_SYSTEM_FIELDS: set[str] = set()

# ── Manual alias overrides for customs-specific abbreviations ──
# These are abbreviations/acronyms that can't be guessed from word-splitting.
_MANUAL_ALIASES: dict[str, list[str]] = {
    # Export template
    "Item_RITC":                ["HSN Code", "HSN/SAC Code", "RITC Code", "Tariff Code",
                                 "CTN"],
    "Item_Desc":                ["Description of Goods", "Item Description",
                                 "Product Description", "Material Description",
                                 "Mat Desc", "Product Name", "Goods Description"],
    "Item_Desc1":               ["Description of Goods", "Item Description",
                                 "Product Description", "Material Description",
                                 "Mat Desc", "Product Name", "Goods Description"],
    "Item_Gen_Desc":            ["General Description"],
    "Item_Ser_No":              ["Serial Number", "SL No", "S.No", "Item Number",
                                 "Serial No"],
    "SL_No":                    ["Serial Number", "S.No", "Item Number", "Item_Ser_No",
                                 "Serial No"],
    "Item_Unit1":               ["Unit of Measure", "UOM", "Unit", "UQC"],
    "Item_Unit2":               ["Secondary Unit", "Second UOM"],
    "Item_Unit":                ["Unit of Measure", "UOM", "Unit", "UQC"],
    "Item_Unit_Price":          ["Unit Price", "Rate", "Price per Unit",
                                 "Rate per Unit"],
    "Item_Qty":                 ["Quantity", "Qty", "Number of Items"],
    "Item_Country_Org":         ["Country of Origin", "Origin Country", "COO"],
    "Item_Scheme_Code":         ["Scheme Code", "Export Scheme Code"],
    "Item_SchemeCode":          ["Scheme Code", "Export Scheme Code"],
    "Itm_Manu_Code":           ["Manufacturer Code"],
    "Itm_Manu_Add1":           ["Manufacturer Address Line 1"],
    "Itm_Manu_Add2":           ["Manufacturer Address Line 2"],
    "Itm_Manu_City":           ["Manufacturer City"],
    "Itm_Manu_Pin":            ["Manufacturer PIN Code"],
    "Itm_Manu_Cntry":          ["Manufacturer Country"],
    "Itm_Manu_CntrySubDiv":    ["Manufacturer State/Province"],
    "Itm_Source_Cntry":        ["Source Country", "Country of Origin",
                                 "Country of Origin of Goods", "Origin Country",
                                 "Made In", "Origin"],
    "Itm_Transit_Cntry":       ["Transit Country", "Transshipment Country",
                                 "Via Country"],
    "ItmIGSTstatus":           ["IGST Status", "LUT", "Bond Status",
                                 "GST Payment Status"],
    "ItmTaxableVal":           ["Taxable Value", "Assessable Value",
                                 "Taxable Amount"],
    "Item_Taxable_Val":        ["Taxable Value", "Assessable Value",
                                 "Taxable Amount", "Extended Price",
                                 "Line Total", "Amount"],
    "ItmIGSTamt":              ["IGST Amount"],
    "ItmIGSTPer":              ["IGST Percentage", "IGST Rate"],
    "Item_PMV":                ["Previous Market Value", "PMV"],
    "Item_End_Use":            ["End Use", "EU Code", "End Use Code"],
    "Item_JNoti_No":           ["Joint Notification Flag"],
    "item_Cess":               ["Cess Flag"],
    "item_accessory":          ["Accessory Flag"],
    "item_thirdparty":         ["Third Party Flag"],
    "item_Quota":              ["Quota Flag"],
    "item_AR4":                ["AR4 Flag"],
    "Item_Reward":             ["Reward Flag", "RoDTEP Flag", "RoDTEP Declaration",
                                 "RoDTEP", "MEIS", "Remission of Duties and Taxes on Exported Products",
                                 "Reward Scheme", "Export Reward"],
    "Item_STR":                ["STR Flag"],
    "Item_BCD_Ntfn":           ["BCD Notification Number",
                                 "Basic Customs Duty Notification"],
    "Item_BCD_Ntfn_SLNo":      ["BCD Notification Serial Number"],
    "Item_CVD_Ntfn":           ["CVD Notification Number"],
    "Item_CVD_Ntfn_SLNo":      ["CVD Notification Serial Number"],
    "PER":                     ["Per Unit", "Price Per"],
    "ItmHawb":                 ["HAWB Number", "House Airway Bill"],
    "ItmTotpkg":               ["Total Packages", "Total Cases",
                                 "No. of Packages", "Number of Packages"],
    "Printed":                 ["Print Status"],
    "FTA Code":                ["Free Trade Agreement Code",
                                 "Preferential Tariff Code", "FTA"],
    "State Code":              ["State Code", "Exporter State Code", "State Code : MH(12)",
                                 "State Code :", "State Code Number", "State", "State / State Code"],
    "District Code":           ["District Code", "District Code Number", "District"],
    "Qty Tariff":              ["Quantity for Tariff", "Tariff Quantity"],
    "Unit Tariff":             ["Unit for Tariff", "Tariff Unit"],
    # Import template extras
    "Item_SWC_Ntfn":           ["SWC Notification"],
    "Item_SWC_Ntfn_SlNo":      ["SWC Notification Serial"],
}


def _auto_alias(col_name: str) -> str:
    """
    Generate a human-readable alias by splitting on underscores, expanding
    common abbreviations, and title-casing.
    """
    _ABBREVS = {
        "ser": "Serial", "no": "Number", "desc": "Description",
        "qty": "Quantity", "org": "Origin", "cntry": "Country",
        "manu": "Manufacturer", "add": "Address", "itm": "Item",
        "sl": "Serial", "ntfn": "Notification", "mfd": "Manufactured",
        "per": "Percentage", "val": "Value", "amt": "Amount",
        "regn": "Registration", "lic": "License",
    }
    parts = col_name.replace("_", " ").split()
    expanded = [_ABBREVS.get(p.lower(), p.capitalize()) for p in parts]
    return " ".join(expanded)


def _infer_type(col_name: str) -> str:
    name = col_name.lower()
    # Force string for known Y/N flags, codes, and text fields
    if name in _STRING_FORCE:
        return "string"
    if any(h in name for h in _DATE_HINTS):
        return "string"  # dates kept as ISO strings, not JSON numbers
    if any(h in name for h in _NUMBER_HINTS):
        return "number"
    return "string"


def load_template_schema(xlsx_path: str | Path) -> dict[str, dict[str, Any]]:
    """
    Returns, per sheet:
        {
          "ITEM": {
             "fields": ["Item_Ser_No", "Item_Desc", ...],
             "types":  {"Item_Ser_No": "string", "Item_Qty": "number", ...},
             "aliases": {"Item_Ser_No": ["Serial Number", "S.No", ...], ...},
             "link_field": "Item_No" | "Item_Ser_No" | None
          },
          ...
        }

    `aliases` maps each column name to a list of common invoice labels
    that the LLM should recognise as that column.

    `link_field` is whichever column ties a sheet's rows back to a
    specific ITEM row (most sheets use "Item_No"; ITEM itself uses
    "Item_Ser_No"), so extracted rows can be joined later.
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    schema: dict[str, dict[str, Any]] = {}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        headers = [c.value for c in ws[1] if c.value is not None]
        headers = [str(h).strip() for h in headers]

        fields = [h for h in headers if h.lower() not in _SYSTEM_FIELDS]
        types = {h: _infer_type(h) for h in fields}

        # Build aliases: manual overrides first, then auto-generate
        aliases: dict[str, list[str]] = {}
        for h in fields:
            if h in _MANUAL_ALIASES:
                aliases[h] = _MANUAL_ALIASES[h]
            else:
                aliases[h] = [_auto_alias(h)]

        link_field = next(
            (h for h in fields if h.replace(" ", "_").lower() in
             ("item_no", "item_ser_no", "sl_no")),
            None,
        )

        schema[sheet_name] = {
            "fields": fields,
            "types": types,
            "aliases": aliases,
            "link_field": link_field,
        }
    return schema


if __name__ == "__main__":
    import json, sys
    path = sys.argv[1] if len(sys.argv) > 1 else "templates/EXP_TEMPLET.xlsx"
    s = load_template_schema(path)
    print(json.dumps(s, indent=2))
